from __future__ import annotations

import math
from collections import Counter
from dataclasses import dataclass
from pathlib import Path

import pandas as pd
from docx import Document
from ortools.constraint_solver import pywrapcp, routing_enums_pb2


VEHICLE_SPEED_KMPH = 50.0
ROUTE_START_HOUR = 6
LAST_DELIVERY_DEADLINE_MINUTES = 4 * 60
END_OF_DAY_MINUTES = 24 * 60
FIXED_COST_INR = 1000.0
RATE_PER_KM_INR = 10.0
RATE_PER_HOUR_INR = 50.0
VARIABLE_COST_PER_KM_INR = RATE_PER_KM_INR + RATE_PER_HOUR_INR / VEHICLE_SPEED_KMPH
COST_SCALE = 100


@dataclass(frozen=True)
class ModelInputs:
    mothers: pd.DataFrame
    daughters: pd.DataFrame
    nodes: list[dict[str, object]]
    center_name_by_node: list[int]
    mother_node_by_center: dict[int, int]
    daughter_nodes: list[int]
    distance_km: list[list[float]]
    time_minutes: list[list[int]]
    cost_units: list[list[int]]
    vehicle_starts: list[int]
    vehicle_ends: list[int]
    vehicle_labels: list[str]


def haversine_km(lat1: float, lon1: float, lat2: float, lon2: float) -> float:
    radius_km = 6371.0
    lat1_rad = math.radians(lat1)
    lat2_rad = math.radians(lat2)
    delta_lat = math.radians(lat2 - lat1)
    delta_lon = math.radians(lon2 - lon1)
    a_value = (
        math.sin(delta_lat / 2) ** 2
        + math.cos(lat1_rad) * math.cos(lat2_rad) * math.sin(delta_lon / 2) ** 2
    )
    return 2 * radius_km * math.asin(math.sqrt(a_value))


def minutes_to_clock(minutes_from_start: int) -> str:
    total_minutes = ROUTE_START_HOUR * 60 + int(minutes_from_start)
    hour, minute = divmod(total_minutes, 60)
    return f"{hour:02d}:{minute:02d}"


def load_inputs(input_file: Path) -> ModelInputs:
    raw_df = pd.read_excel(input_file)
    raw_df["center_name"] = raw_df["center_name"].astype(int)

    mothers = (
        raw_df[raw_df["Category"].str.lower() == "mother"]
        .sort_values("center_name")
        .reset_index(drop=True)
    )
    daughters = (
        raw_df[raw_df["Category"].str.lower() == "daughter"]
        .sort_values("center_name")
        .reset_index(drop=True)
    )

    nodes: list[dict[str, object]] = []
    center_name_by_node: list[int] = []
    mother_node_by_center: dict[int, int] = {}

    for row in mothers.itertuples(index=False):
        node_index = len(nodes)
        center_name = int(row.center_name)
        nodes.append(
            {
                "center_name": center_name,
                "latitude": float(row.latitude),
                "longitude": float(row.longitude),
                "category": "Mother",
            }
        )
        center_name_by_node.append(center_name)
        mother_node_by_center[center_name] = node_index

    daughter_nodes: list[int] = []
    for row in daughters.itertuples(index=False):
        node_index = len(nodes)
        center_name = int(row.center_name)
        nodes.append(
            {
                "center_name": center_name,
                "latitude": float(row.latitude),
                "longitude": float(row.longitude),
                "category": "Daughter",
            }
        )
        center_name_by_node.append(center_name)
        daughter_nodes.append(node_index)

    coordinates = [(float(node["latitude"]), float(node["longitude"])) for node in nodes]
    node_count = len(nodes)
    distance_km = [[0.0 for _ in range(node_count)] for _ in range(node_count)]
    time_minutes = [[0 for _ in range(node_count)] for _ in range(node_count)]
    cost_units = [[0 for _ in range(node_count)] for _ in range(node_count)]

    for from_node in range(node_count):
        from_lat, from_lon = coordinates[from_node]
        for to_node in range(node_count):
            if from_node == to_node:
                continue
            to_lat, to_lon = coordinates[to_node]
            distance = haversine_km(from_lat, from_lon, to_lat, to_lon)
            distance_km[from_node][to_node] = distance
            travel_minutes = math.ceil(distance * 60 / VEHICLE_SPEED_KMPH)
            time_minutes[from_node][to_node] = travel_minutes
            cost_units[from_node][to_node] = int(round(distance * VARIABLE_COST_PER_KM_INR * COST_SCALE))

    nearest_vehicle_counts: Counter[int] = Counter()
    for daughter_node in daughter_nodes:
        nearest_center = min(
            mother_node_by_center,
            key=lambda center_name: distance_km[mother_node_by_center[center_name]][daughter_node],
        )
        nearest_vehicle_counts[nearest_center] += 1

    vehicle_starts: list[int] = []
    vehicle_ends: list[int] = []
    vehicle_labels: list[str] = []
    for mother_center in sorted(mother_node_by_center):
        vehicle_count = max(1, nearest_vehicle_counts.get(mother_center, 0))
        mother_node = mother_node_by_center[mother_center]
        for sequence in range(1, vehicle_count + 1):
            vehicle_starts.append(mother_node)
            vehicle_ends.append(mother_node)
            vehicle_labels.append(f"M{mother_center}_V{sequence:02d}")

    return ModelInputs(
        mothers=mothers,
        daughters=daughters,
        nodes=nodes,
        center_name_by_node=center_name_by_node,
        mother_node_by_center=mother_node_by_center,
        daughter_nodes=daughter_nodes,
        distance_km=distance_km,
        time_minutes=time_minutes,
        cost_units=cost_units,
        vehicle_starts=vehicle_starts,
        vehicle_ends=vehicle_ends,
        vehicle_labels=vehicle_labels,
    )


def solve_model(model_inputs: ModelInputs) -> tuple[pywrapcp.RoutingIndexManager, pywrapcp.RoutingModel, pywrapcp.Assignment, pywrapcp.RoutingDimension]:
    manager = pywrapcp.RoutingIndexManager(
        len(model_inputs.nodes),
        len(model_inputs.vehicle_starts),
        model_inputs.vehicle_starts,
        model_inputs.vehicle_ends,
    )
    routing = pywrapcp.RoutingModel(manager)

    def cost_callback(from_index: int, to_index: int) -> int:
        from_node = manager.IndexToNode(from_index)
        to_node = manager.IndexToNode(to_index)
        return model_inputs.cost_units[from_node][to_node]

    def time_callback(from_index: int, to_index: int) -> int:
        from_node = manager.IndexToNode(from_index)
        to_node = manager.IndexToNode(to_index)
        return model_inputs.time_minutes[from_node][to_node]

    cost_callback_index = routing.RegisterTransitCallback(cost_callback)
    routing.SetArcCostEvaluatorOfAllVehicles(cost_callback_index)

    time_callback_index = routing.RegisterTransitCallback(time_callback)
    routing.AddDimension(
        time_callback_index,
        0,
        END_OF_DAY_MINUTES,
        True,
        "Time",
    )
    time_dimension = routing.GetDimensionOrDie("Time")

    for daughter_node in model_inputs.daughter_nodes:
        node_index = manager.NodeToIndex(daughter_node)
        time_dimension.CumulVar(node_index).SetRange(0, LAST_DELIVERY_DEADLINE_MINUTES)

    fixed_cost_units = int(round(FIXED_COST_INR * COST_SCALE))
    for vehicle_id in range(len(model_inputs.vehicle_starts)):
        routing.SetFixedCostOfVehicle(fixed_cost_units, vehicle_id)
        time_dimension.CumulVar(routing.Start(vehicle_id)).SetRange(0, 0)
        time_dimension.CumulVar(routing.End(vehicle_id)).SetRange(0, END_OF_DAY_MINUTES)
        routing.AddVariableMinimizedByFinalizer(time_dimension.CumulVar(routing.End(vehicle_id)))

    search_parameters = pywrapcp.DefaultRoutingSearchParameters()
    search_parameters.first_solution_strategy = routing_enums_pb2.FirstSolutionStrategy.PARALLEL_CHEAPEST_INSERTION
    search_parameters.local_search_metaheuristic = routing_enums_pb2.LocalSearchMetaheuristic.GUIDED_LOCAL_SEARCH
    search_parameters.time_limit.FromSeconds(90)
    search_parameters.log_search = False

    solution = routing.SolveWithParameters(search_parameters)
    if solution is None:
        raise RuntimeError("OR-Tools could not find a feasible routing solution.")

    return manager, routing, solution, time_dimension


def build_outputs(
    model_inputs: ModelInputs,
    manager: pywrapcp.RoutingIndexManager,
    routing: pywrapcp.RoutingModel,
    solution: pywrapcp.Assignment,
    time_dimension: pywrapcp.RoutingDimension,
) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame, dict[str, float]]:
    assignment_rows: list[dict[str, object]] = []
    route_rows: list[dict[str, object]] = []

    for vehicle_id, vehicle_label in enumerate(model_inputs.vehicle_labels):
        if not routing.IsVehicleUsed(solution, vehicle_id):
            continue

        start_index = routing.Start(vehicle_id)
        start_node = manager.IndexToNode(start_index)
        mother_center = model_inputs.center_name_by_node[start_node]
        route_id = f"{mother_center}_{vehicle_label}"
        route_distance_km = 0.0
        route_path = [str(mother_center)]
        sequence = 0
        customer_count = 0
        last_delivery_minutes = 0
        index = start_index

        while not routing.IsEnd(index):
            next_index = solution.Value(routing.NextVar(index))
            from_node = manager.IndexToNode(index)
            to_node = manager.IndexToNode(next_index)
            route_distance_km += model_inputs.distance_km[from_node][to_node]

            if not routing.IsEnd(next_index):
                sequence += 1
                customer_count += 1
                arrival_minutes = solution.Value(time_dimension.CumulVar(next_index))
                last_delivery_minutes = arrival_minutes
                daughter_center = model_inputs.center_name_by_node[to_node]
                route_path.append(str(daughter_center))
                assignment_rows.append(
                    {
                        "daughter_center": daughter_center,
                        "assigned_mother": mother_center,
                        "route_id": route_id,
                        "vehicle_label": vehicle_label,
                        "stop_sequence": sequence,
                        "arrival_minutes_from_6am": arrival_minutes,
                        "arrival_clock": minutes_to_clock(arrival_minutes),
                    }
                )

            index = next_index

        route_path.append(str(mother_center))
        route_duration_minutes = solution.Value(time_dimension.CumulVar(routing.End(vehicle_id)))
        variable_cost_inr = route_distance_km * VARIABLE_COST_PER_KM_INR
        total_cost_inr = FIXED_COST_INR + variable_cost_inr
        route_rows.append(
            {
                "route_id": route_id,
                "vehicle_label": vehicle_label,
                "mother_center": mother_center,
                "stops": customer_count,
                "path": " -> ".join(route_path),
                "last_delivery_minutes_from_6am": last_delivery_minutes,
                "last_delivery_clock": minutes_to_clock(last_delivery_minutes),
                "returned_to_mother_clock": minutes_to_clock(route_duration_minutes),
                "route_distance_km": round(route_distance_km, 2),
                "route_drive_hours": round(route_duration_minutes / 60, 2),
                "fixed_cost_inr": FIXED_COST_INR,
                "variable_cost_inr": round(variable_cost_inr, 2),
                "total_cost_inr": round(total_cost_inr, 2),
            }
        )

    assignments_df = pd.DataFrame(assignment_rows).sort_values(
        ["assigned_mother", "route_id", "stop_sequence"]
    )
    routes_df = pd.DataFrame(route_rows).sort_values(["mother_center", "route_id"])

    if assignments_df["daughter_center"].nunique() != len(model_inputs.daughters):
        raise RuntimeError("The solution did not assign every daughter center exactly once.")

    mother_summary_df = (
        assignments_df.groupby("assigned_mother", as_index=False)
        .agg(
            assigned_daughters=("daughter_center", "count"),
            routes_used=("route_id", "nunique"),
        )
        .sort_values("assigned_mother")
    )

    total_distance_km = float(routes_df["route_distance_km"].sum())
    total_fixed_cost_inr = float(routes_df["fixed_cost_inr"].sum())
    total_variable_cost_inr = float(routes_df["variable_cost_inr"].sum())
    total_cost_inr = float(routes_df["total_cost_inr"].sum())

    summary_values = {
        "vehicle_speed_kmph": VEHICLE_SPEED_KMPH,
        "mother_centers": float(len(model_inputs.mothers)),
        "daughter_centers": float(len(model_inputs.daughters)),
        "routes_used": float(len(routes_df)),
        "total_distance_km": round(total_distance_km, 2),
        "total_fixed_cost_inr": round(total_fixed_cost_inr, 2),
        "total_variable_cost_inr": round(total_variable_cost_inr, 2),
        "total_cost_inr": round(total_cost_inr, 2),
    }
    summary_df = pd.DataFrame(
        [{"metric": metric, "value": value} for metric, value in summary_values.items()]
    )

    return assignments_df, routes_df, mother_summary_df, summary_df, summary_values


def write_excel(
    output_file: Path,
    assignments_df: pd.DataFrame,
    routes_df: pd.DataFrame,
    mother_summary_df: pd.DataFrame,
    summary_df: pd.DataFrame,
) -> None:
    with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
        assignments_df.to_excel(writer, sheet_name="Assignments", index=False)
        routes_df.to_excel(writer, sheet_name="Routes", index=False)
        mother_summary_df.to_excel(writer, sheet_name="MotherSummary", index=False)
        summary_df.to_excel(writer, sheet_name="Summary", index=False)


def write_methodology_docx(
    output_file: Path,
    summary_values: dict[str, float],
    mother_summary_df: pd.DataFrame,
) -> None:
    document = Document()
    document.add_heading("Mother-Daughter Routing Methodology", level=1)

    document.add_paragraph(
        "This solution models the assignment as a multi-depot vehicle routing problem solved with Google OR-Tools. "
        "Each route starts from a mother center at 6:00 AM, visits one or more daughter centers, reaches the last daughter by 10:00 AM, and then returns to the same mother center."
    )

    document.add_heading("Model Setup", level=2)
    document.add_paragraph(
        "Distance is computed with haversine geometry because only latitude and longitude were provided. "
        "Vehicle speed is fixed at 50 km/h, so travel time is distance divided by 50. "
        "The commercial objective is fixed cost plus variable cost, where variable cost equals Rs. 10 per km and Rs. 50 per hour. "
        "At 50 km/h this becomes Rs. 11 per km over the full round trip."
    )

    document.add_heading("Assumptions", level=2)
    for assumption in [
        "A homogeneous Tata Ace fleet is available across mother centers up to the modeled vehicle pool.",
        "No capacity limit or unloading time was specified, so service time at daughter centers is assumed to be zero.",
        "Every daughter center must be assigned to exactly one mother center and served by exactly one route.",
        "The 10:00 AM cutoff applies to the last daughter visit, while the return to the mother center can happen later.",
    ]:
        document.add_paragraph(assumption, style="List Bullet")

    document.add_heading("Results", level=2)
    document.add_paragraph(
        "The optimized plan serves all daughter centers with "
        f"{int(summary_values['routes_used'])} routes, covering {summary_values['total_distance_km']:.2f} km at a total cost of Rs. {summary_values['total_cost_inr']:.2f}."
    )

    table = document.add_table(rows=1, cols=3)
    header = table.rows[0].cells
    header[0].text = "Mother Center"
    header[1].text = "Assigned Daughters"
    header[2].text = "Routes Used"
    for row in mother_summary_df.itertuples(index=False):
        cells = table.add_row().cells
        cells[0].text = str(int(row.assigned_mother))
        cells[1].text = str(int(row.assigned_daughters))
        cells[2].text = str(int(row.routes_used))

    document.save(output_file)


def write_summary_markdown(
    output_file: Path,
    summary_values: dict[str, float],
    mother_summary_df: pd.DataFrame,
    routes_df: pd.DataFrame,
) -> None:
    lines = [
        "# OR-Tools Routing Summary",
        "",
        f"- Vehicle speed: {VEHICLE_SPEED_KMPH:.0f} km/h",
        f"- Mother centers: {int(summary_values['mother_centers'])}",
        f"- Daughter centers: {int(summary_values['daughter_centers'])}",
        f"- Routes used: {int(summary_values['routes_used'])}",
        f"- Total distance: {summary_values['total_distance_km']:.2f} km",
        f"- Total fixed cost: Rs. {summary_values['total_fixed_cost_inr']:.2f}",
        f"- Total variable cost: Rs. {summary_values['total_variable_cost_inr']:.2f}",
        f"- Total cost: Rs. {summary_values['total_cost_inr']:.2f}",
        "",
        "## Mother Assignment Summary",
        "",
        "| Mother Center | Assigned Daughters | Routes Used |",
        "| --- | ---: | ---: |",
    ]
    for row in mother_summary_df.itertuples(index=False):
        lines.append(
            f"| {int(row.assigned_mother)} | {int(row.assigned_daughters)} | {int(row.routes_used)} |"
        )

    lines.extend(
        [
            "",
            "## Longest Routes",
            "",
            "| Route ID | Mother Center | Stops | Distance (km) | Last Delivery | Total Cost (Rs.) |",
            "| --- | ---: | ---: | ---: | --- | ---: |",
        ]
    )
    longest_routes = routes_df.sort_values("route_distance_km", ascending=False).head(10)
    for row in longest_routes.itertuples(index=False):
        lines.append(
            f"| {row.route_id} | {int(row.mother_center)} | {int(row.stops)} | {row.route_distance_km:.2f} | {row.last_delivery_clock} | {row.total_cost_inr:.2f} |"
        )

    output_file.write_text("\n".join(lines), encoding="utf-8")


def main() -> None:
    workspace_dir = Path(__file__).resolve().parent
    input_file = workspace_dir / "data_1.xlsx"
    output_dir = workspace_dir / "outputs"
    output_dir.mkdir(exist_ok=True)

    model_inputs = load_inputs(input_file)
    manager, routing, solution, time_dimension = solve_model(model_inputs)
    assignments_df, routes_df, mother_summary_df, summary_df, summary_values = build_outputs(
        model_inputs,
        manager,
        routing,
        solution,
        time_dimension,
    )

    excel_file = output_dir / "or_tools_solution.xlsx"
    docx_file = output_dir / "or_tools_methodology.docx"
    summary_file = output_dir / "or_tools_summary.md"

    write_excel(excel_file, assignments_df, routes_df, mother_summary_df, summary_df)
    write_methodology_docx(docx_file, summary_values, mother_summary_df)
    write_summary_markdown(summary_file, summary_values, mother_summary_df, routes_df)

    print(f"Solved assignment with vehicle speed fixed at {VEHICLE_SPEED_KMPH:.0f} km/h")
    print(f"Routes used: {int(summary_values['routes_used'])}")
    print(f"Total distance (km): {summary_values['total_distance_km']:.2f}")
    print(f"Total cost (Rs.): {summary_values['total_cost_inr']:.2f}")
    print(f"Excel output: {excel_file}")
    print(f"Methodology doc: {docx_file}")
    print(f"Summary: {summary_file}")


if __name__ == "__main__":
    main()